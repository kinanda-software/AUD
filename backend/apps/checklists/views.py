from django.db import transaction
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from .models import (
    ChecklistItem,
    ChecklistResponse,
    ChecklistTemplate,
    EngagementChecklist,
)
from .serializers import (
    BulkResponseSerializer,
    ChecklistItemSerializer,
    ChecklistResponseSerializer,
    ChecklistTemplateSerializer,
    EngagementChecklistSerializer,
)


class ChecklistTemplateViewSet(viewsets.ModelViewSet):
    queryset = (
        ChecklistTemplate.objects
        .select_related("created_by")
        .prefetch_related("items")
        .all()
    )

    serializer_class = ChecklistTemplateSerializer

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def clone(self, request, pk=None):
        """Deep-copy a template together with its items."""
        source = self.get_object()

        clone = ChecklistTemplate.objects.create(
            name=f"{source.name} (Copy)",
            description=source.description,
            category=source.category,
            created_by=request.user,
        )

        id_map = {}
        for item in source.items.all():
            new_item = ChecklistItem.objects.create(
                template=clone,
                order=item.order,
                question=item.question,
                response_type=item.response_type,
                weight=item.weight,
                condition_value=item.condition_value,
            )
            id_map[item.pk] = new_item

        for old_id, new_item in id_map.items():
            old_parent_id = (
                source.items.get(pk=old_id).parent_id
            )
            if old_parent_id and old_parent_id in id_map:
                new_item.parent = id_map[old_parent_id]
                new_item.save(update_fields=["parent"])

        serializer = self.get_serializer(clone)
        return Response(
            serializer.data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=False,
        methods=["post"],
        parser_classes=(MultiPartParser, FormParser),
    )
    @transaction.atomic
    def import_excel(self, request):
        """
        Create a template from an .xlsx workbook.

        Expected columns (header row):
            question (required), response_type, weight,
            parent, condition_value, order

        `parent` holds the exact question text of the parent
        row, enabling conditional sub-questions.
        """
        upload = request.FILES.get("file")
        if upload is None:
            raise ValidationError({"file": "Attach an .xlsx workbook."})

        name = (request.data.get("name") or "").strip()
        if not name:
            raise ValidationError({"name": "Provide a template name."})

        if not upload.name.lower().endswith(".xlsx"):
            raise ValidationError(
                {"file": "Only .xlsx workbooks are supported."}
            )

        try:
            from openpyxl import load_workbook
            workbook = load_workbook(
                upload, read_only=True, data_only=True
            )
        except Exception:
            raise ValidationError(
                {"file": "Could not read this Excel workbook."}
            )

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()

        if not rows:
            raise ValidationError(
                {"file": "The workbook has no rows."}
            )

        header = [
            str(cell).strip().lower() if cell else ""
            for cell in rows[0]
        ]
        if "question" not in header:
            raise ValidationError(
                {"file": "The first row must include a 'question' column."}
            )

        def column(row, key, default=""):
            if key not in header:
                return default
            value = row[header.index(key)]
            return value if value is not None else default

        valid_types = {
            value for value, _ in
            ChecklistItem.ResponseType.choices
        }

        template = ChecklistTemplate.objects.create(
            name=name,
            category=(
                str(request.data.get("category") or "General")
            ),
            description=str(
                request.data.get("description") or ""
            ),
            created_by=request.user,
        )

        # First pass: create items without parents.
        created = {}
        pending_parents = []
        for index, row in enumerate(rows[1:], start=1):
            question = str(column(row, "question")).strip()
            if not question:
                continue

            response_type = (
                str(column(row, "response_type"))
                .strip().lower() or "yes_no_na"
            )
            if response_type not in valid_types:
                response_type = "yes_no_na"

            try:
                weight = float(column(row, "weight", 1) or 1)
            except (TypeError, ValueError):
                weight = 1.0

            try:
                order = int(column(row, "order", index) or index)
            except (TypeError, ValueError):
                order = index

            item = ChecklistItem.objects.create(
                template=template,
                order=order,
                question=question,
                response_type=response_type,
                weight=weight,
                condition_value=str(
                    column(row, "condition_value")
                ).strip().lower(),
            )
            created[question] = item

            parent_question = str(
                column(row, "parent")
            ).strip()
            if parent_question:
                pending_parents.append((item, parent_question))

        # Second pass: link conditional sub-questions.
        for item, parent_question in pending_parents:
            parent = created.get(parent_question)
            if parent is not None and parent.pk != item.pk:
                item.parent = parent
                item.save(update_fields=["parent"])

        serializer = self.get_serializer(template)
        return Response(
            serializer.data,
            status=status.HTTP_201_CREATED,
        )


class ChecklistItemViewSet(viewsets.ModelViewSet):
    queryset = (
        ChecklistItem.objects
        .select_related("template", "parent")
        .all()
    )
    serializer_class = ChecklistItemSerializer

    def get_queryset(self):
        queryset = super().get_queryset()
        template_id = self.request.query_params.get("template")
        if template_id:
            queryset = queryset.filter(template_id=template_id)
        return queryset


class EngagementChecklistViewSet(viewsets.ModelViewSet):
    queryset = (
        EngagementChecklist.objects
        .select_related(
            "engagement",
            "template",
            "created_by",
        )
        .prefetch_related("responses", "responses__item")
        .all()
    )

    serializer_class = EngagementChecklistSerializer

    def get_queryset(self):
        queryset = super().get_queryset()

        engagement_id = self.request.query_params.get(
            "engagement"
        )
        if engagement_id:
            queryset = queryset.filter(
                engagement_id=engagement_id
            )

        return queryset

    @transaction.atomic
    def perform_create(self, serializer):
        checklist = serializer.save(
            created_by=self.request.user
        )
        # Pre-create blank responses for every template item.
        ChecklistResponse.objects.bulk_create([
            ChecklistResponse(
                engagement_checklist=checklist,
                item=item,
            )
            for item in checklist.template.items.all()
        ])

    @action(detail=True, methods=["get"])
    def score(self, request, pk=None):
        checklist = self.get_object()
        return Response(checklist.compute_score())

    @action(
        detail=True,
        methods=["get", "put"],
        url_path="responses",
    )
    @transaction.atomic
    def responses(self, request, pk=None):
        checklist = self.get_object()

        if request.method == "GET":
            serializer = ChecklistResponseSerializer(
                checklist.responses.all(),
                many=True,
            )
            return Response(serializer.data)

        serializer = BulkResponseSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        valid_item_ids = set(
            checklist.template.items.values_list(
                "pk", flat=True
            )
        )

        for entry in serializer.validated_data["responses"]:
            item_id = entry["item"]
            if item_id not in valid_item_ids:
                raise ValidationError(
                    f"Item {item_id} does not belong to this "
                    "checklist's template."
                )
            ChecklistResponse.objects.update_or_create(
                engagement_checklist=checklist,
                item_id=item_id,
                defaults={
                    "value": entry.get("value", ""),
                    "rating": entry.get("rating"),
                    "text_value": entry.get("text_value", ""),
                    "comment": entry.get("comment", ""),
                },
            )

        checklist.status = (
            EngagementChecklist.Status.COMPLETED
            if serializer.validated_data["complete"]
            else EngagementChecklist.Status.IN_PROGRESS
        )
        checklist.save(update_fields=["status", "updated_at"])

        return Response(
            EngagementChecklistSerializer(
                checklist,
                context={"request": request},
            ).data
        )
