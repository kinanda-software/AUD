import csv
import hashlib
import io
import json
from zipfile import BadZipFile
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

from .accounting_controls import require_manager
from apps.engagements.models import Engagement
from .models import (
    ChartOfAccount,
    FinancialAccountMapping,
    FinancialDataImport,
    TrialBalance,
    TrialBalanceLine,
)
from .workflow_services import record_financial_event


MAX_IMPORT_SIZE = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10000
REQUIRED_COLUMNS = ("account_code", "debit", "credit")
ZERO = Decimal("0.00")


def _json_safe(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool, Decimal)):
        return str(value).strip()
    return str(value).strip()


def _read_upload(upload):
    if upload.size > MAX_IMPORT_SIZE:
        raise ValidationError({"file": "Files must be no larger than 10 MB."})
    filename = upload.name or "financial-data"
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    content = upload.read()
    if extension == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValidationError({"file": "CSV files must use UTF-8 encoding."}) from error
        reader = csv.DictReader(io.StringIO(text, newline=""))
        headers = reader.fieldnames or []
        if not headers or len(headers) != len(set(headers)):
            raise ValidationError({"file": "The CSV must have a header row with unique column names."})
        rows = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_IMPORT_ROWS:
                raise ValidationError({"file": f"Files may contain at most {MAX_IMPORT_ROWS} data rows."})
        source_format = "csv"
    elif extension == "xlsx":
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet is None:
                raise ValidationError({"file": "The XLSX workbook has no active worksheet."})
            values = sheet.iter_rows(values_only=True)
            headers = [_json_safe(value) for value in next(values, ())]
            if not headers or len(headers) != len(set(headers)):
                raise ValidationError({"file": "The first worksheet must have unique column headings."})
            rows = []
            for values_row in values:
                if any(value is not None for value in values_row):
                    rows.append(dict(zip(headers, (_json_safe(value) for value in values_row))))
                    if len(rows) > MAX_IMPORT_ROWS:
                        raise ValidationError({"file": f"Files may contain at most {MAX_IMPORT_ROWS} data rows."})
        except ValidationError:
            raise
        except (BadZipFile, EOFError, InvalidFileException, KeyError, OSError, SyntaxError, ValueError) as error:
            raise ValidationError({"file": "The XLSX workbook could not be read."}) from error
        finally:
            if workbook is not None:
                workbook.close()
        source_format = "xlsx"
    else:
        raise ValidationError({"file": "Upload a CSV or XLSX file."})
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValidationError({"file": f"Files may contain at most {MAX_IMPORT_ROWS} data rows."})
    return filename, source_format, content, headers, rows


def _resolve_account(engagement_id, source_system, external_code):
    mapping = FinancialAccountMapping.objects.filter(
        engagement_id=engagement_id,
        source_system__iexact=source_system,
        external_code=external_code,
    ).select_related("account").first()
    account = mapping.account if mapping else ChartOfAccount.objects.filter(
        engagement_id=engagement_id,
        account_code=external_code,
    ).first()
    if account is None:
        return None, "Account code is not mapped to this engagement."
    if account.engagement_id != engagement_id:
        return None, "Mapped account belongs to another engagement."
    if not account.is_active:
        return None, "Mapped account is inactive."
    return account, None


def create_data_import(data, actor):
    upload = data.get("file")
    if upload:
        filename, source_format, content, headers, input_rows = _read_upload(upload)
        source_fingerprint = hashlib.sha256(content).hexdigest()
    else:
        filename = "API payload"
        source_format = "json"
        input_rows = data["rows"]
        headers = list(input_rows[0]) if input_rows else []
        source_fingerprint = hashlib.sha256(
            json.dumps(input_rows, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
        ).hexdigest()

    if not input_rows:
        raise ValidationError({"file": "The source contains no data rows."})

    source_system = data["source_system"].strip()
    column_mapping = data.get("column_mapping") or {
        field: field for field in REQUIRED_COLUMNS
    }
    missing_fields = sorted(set(REQUIRED_COLUMNS) - set(column_mapping))
    if missing_fields:
        raise ValidationError({"column_mapping": f"Map these required fields: {', '.join(missing_fields)}."})
    missing_headers = sorted(
        {column_mapping[field] for field in REQUIRED_COLUMNS} - set(headers)
    )
    if missing_headers:
        raise ValidationError({"column_mapping": f"Source columns not found: {', '.join(missing_headers)}."})

    errors = []
    existing_source = FinancialDataImport.objects.filter(
        engagement=data["engagement"],
        source_system__iexact=source_system,
        source_fingerprint=source_fingerprint,
    ).first()
    if existing_source:
        errors.append({
            "row": None,
            "field": "file",
            "message": f"This source was already received as import #{existing_source.pk}.",
        })

    rows = []
    debit_total = ZERO
    credit_total = ZERO
    seen_accounts = {}
    for row_number, input_row in enumerate(input_rows, start=2 if upload else 1):
        external_code = _json_safe(input_row.get(column_mapping["account_code"]))
        normalized = {
            "source_row": row_number,
            "external_code": external_code,
            "external_name": _json_safe(input_row.get(column_mapping.get("account_name", ""))),
            "account_id": None,
            "account_code": "",
            "account_name": "",
            "debit": "0.00",
            "credit": "0.00",
        }
        row_is_valid = True
        amounts_are_valid = True

        if not external_code:
            errors.append({"row": row_number, "field": "account_code", "message": "Account code is required."})
            row_is_valid = False
        else:
            account, account_error = _resolve_account(data["engagement"].pk, source_system, external_code)
            if account_error:
                errors.append({"row": row_number, "field": "account_code", "message": account_error})
                row_is_valid = False
            else:
                normalized.update({
                    "account_id": account.pk,
                    "account_code": account.account_code,
                    "account_name": account.account_name,
                })
                if account.pk in seen_accounts:
                    errors.append({
                        "row": row_number,
                        "field": "account_code",
                        "message": f"Duplicate mapped account; first seen on row {seen_accounts[account.pk]}.",
                    })
                    row_is_valid = False
                else:
                    seen_accounts[account.pk] = row_number

        amounts = {}
        for field in ("debit", "credit"):
            source_value = _json_safe(input_row.get(column_mapping[field]))
            try:
                amount = Decimal(source_value or "0")
                if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
                    raise InvalidOperation
                amount = amount.quantize(Decimal("0.01"))
                if amount.adjusted() > 15:
                    raise InvalidOperation
                amounts[field] = amount
                normalized[field] = format(amount, ".2f")
            except (InvalidOperation, ValueError):
                errors.append({
                    "row": row_number,
                    "field": field,
                    "message": "Enter a non-negative amount with no more than two decimal places.",
                })
                row_is_valid = False
                amounts_are_valid = False
                amounts[field] = ZERO

        if amounts["debit"] and amounts["credit"]:
            errors.append({"row": row_number, "field": "amount", "message": "A row cannot contain both debit and credit."})
            row_is_valid = False
        elif not amounts["debit"] and not amounts["credit"]:
            errors.append({"row": row_number, "field": "amount", "message": "A row must contain a non-zero debit or credit."})
            row_is_valid = False

        if amounts_are_valid:
            debit_total += amounts["debit"]
            credit_total += amounts["credit"]
        rows.append(normalized)

    if debit_total != credit_total:
        errors.append({
            "row": None,
            "field": "reconciliation",
            "message": f"Debits ({debit_total:.2f}) do not equal credits ({credit_total:.2f}).",
        })
    if TrialBalance.objects.filter(
        engagement=data["engagement"],
        period_start=data["period_start"],
        period_end=data["period_end"],
        currency=data["currency"].upper(),
    ).exists():
        errors.append({
            "row": None,
            "field": "period",
            "message": "A trial balance already exists for this engagement, period, and currency.",
        })

    imported = FinancialDataImport.objects.create(
        engagement=data["engagement"],
        source_system=source_system,
        source_name=filename[:255],
        source_format=source_format,
        source_fingerprint=source_fingerprint,
        period_start=data["period_start"],
        period_end=data["period_end"],
        currency=data["currency"].upper(),
        column_mapping=column_mapping,
        rows=rows,
        validation_errors=errors,
        row_count=len(rows),
        total_debit=debit_total,
        total_credit=credit_total,
        status=FinancialDataImport.Status.VALIDATED if not errors else FinancialDataImport.Status.NEEDS_REVIEW,
        created_by=actor,
    )
    record_financial_event(
        engagement_id=imported.engagement_id,
        actor=actor,
        action="created",
        object_type="financial_data_import",
        object_id=imported.pk,
        details={
            "source_system": imported.source_system,
            "source_name": imported.source_name,
            "source_format": imported.source_format,
            "source_fingerprint": imported.source_fingerprint,
            "row_count": imported.row_count,
            "total_debit": format(imported.total_debit, ".2f"),
            "total_credit": format(imported.total_credit, ".2f"),
            "validation_error_count": len(errors),
            "status": imported.status,
        },
    )
    return imported


@transaction.atomic
def commit_data_import(imported, actor):
    require_manager(actor)
    Engagement.objects.select_for_update().get(pk=imported.engagement_id)
    if imported.status != FinancialDataImport.Status.VALIDATED:
        raise ValidationError({"status": "Only a fully validated import can be committed."})
    if imported.trial_balance_id:
        raise ValidationError({"status": "This import has already been committed."})

    accounts = list(ChartOfAccount.objects.filter(
        engagement_id=imported.engagement_id,
        is_active=True,
        pk__in=[row["account_id"] for row in imported.rows],
    ))
    account_by_id = {account.pk: account for account in accounts}
    if len(account_by_id) != len({row["account_id"] for row in imported.rows}):
        raise ValidationError({"rows": "An account mapping changed or became inactive. Create a new import after correcting it."})
    if TrialBalance.objects.filter(
        engagement_id=imported.engagement_id,
        period_start=imported.period_start,
        period_end=imported.period_end,
        currency=imported.currency,
    ).exists():
        raise ValidationError({"period": "A trial balance already exists for this engagement, period, and currency."})

    trial_balance = TrialBalance.objects.create(
        engagement_id=imported.engagement_id,
        period_start=imported.period_start,
        period_end=imported.period_end,
        currency=imported.currency,
        status=TrialBalance.Status.IMPORTED,
        description=f"Imported from {imported.source_system}: {imported.source_name} (import #{imported.pk})",
    )
    TrialBalanceLine.objects.bulk_create([
        TrialBalanceLine(
            trial_balance=trial_balance,
            account=account_by_id[row["account_id"]],
            account_code=row["account_code"],
            account_name=row["account_name"],
            debit=Decimal(row["debit"]),
            credit=Decimal(row["credit"]),
        )
        for row in imported.rows
    ])
    imported.status = FinancialDataImport.Status.IMPORTED
    imported.trial_balance = trial_balance
    imported.imported_at = timezone.now()
    imported.save(update_fields=["status", "trial_balance", "imported_at"])
    record_financial_event(
        engagement_id=imported.engagement_id,
        actor=actor,
        action="imported",
        object_type="financial_data_import",
        object_id=imported.pk,
        details={
            "trial_balance_id": trial_balance.pk,
            "row_count": imported.row_count,
            "source_fingerprint": imported.source_fingerprint,
        },
    )
    return trial_balance
